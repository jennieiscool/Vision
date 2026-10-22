"""Format-aware edits saved as new copies. Original documents stay intact."""
import asyncio
import copy
import hashlib
import io
import json
import re
import secrets
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import quoteattr
from pathlib import Path

from fastapi import HTTPException, Request
from .workspace_files import SHEETS, spreadsheet, presentation, zip_bound, allowed_file

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
T = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
O = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
ODTEXT = 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'
IMAGE_EDIT = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff', '.gif', '.ico', '.avif', '.heic', '.heif'}
DOCUMENT_EDIT = {'.docx', '.odt', '.pptx'}


def image_plugins(path):
    if path.suffix.lower() in {'.heic', '.heif'}:
        import pillow_heif
        pillow_heif.register_heif_opener()


def fingerprint(path):
    result = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            result.update(chunk)
    return result.hexdigest()


def xml_part(path, slide=0):
    extension = path.suffix.lower()
    part, namespace = ('word/document.xml', W) if extension == '.docx' else ('content.xml', ODTEXT)
    with zipfile.ZipFile(path) as archive:
        if extension == '.pptx':
            # Use the same ordered slide parts as the presentation preview.
            package = ET.fromstring(archive.read('ppt/presentation.xml'))
            from .workspace_files import relations
            rels = relations(archive, 'ppt/presentation.xml')
            slides = package.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}sldId')
            if not 0 <= slide < len(slides):
                raise ValueError('Choose an existing slide.')
            part = rels[slides[slide].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
            namespace = A
        document = ET.fromstring(archive.read(part))
    if extension == '.odt':
        paragraphs = [p for p in document.iter() if p.tag in {'{' + ODTEXT + '}p', '{' + ODTEXT + '}h'}]
    else:
        paragraphs = list(document.iter('{' + namespace + '}p'))
    return part, document, paragraphs, namespace


def paragraph_text(paragraph, namespace):
    if namespace == ODTEXT:
        return ''.join(paragraph.itertext())
    return ''.join(node.text or '' for node in paragraph.iter('{' + namespace + '}t'))


def rewrite_zip(source, destination, changes):
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(destination, 'w') as output:
        for info in original.infolist():
            output.writestr(info, changes.get(info.filename, original.read(info.filename)))


def serialize_xml(document, original):
    # Office uses namespace prefixes inside attribute values (mc:Ignorable).
    # Preserve even unused declarations rather than producing a damaged package.
    namespaces = dict(value for _, value in ET.iterparse(io.BytesIO(original), events=('start-ns',)))
    for prefix, uri in namespaces.items():
        if not re.fullmatch(r'ns\d+', prefix):
            ET.register_namespace(prefix, uri)
    output = ET.tostring(document, encoding='unicode')
    missing = ''
    header = output[:output.index('>')]
    for prefix, uri in namespaces.items():
        attribute = 'xmlns:' + prefix if prefix else 'xmlns'
        if not re.search(r'\s' + re.escape(attribute) + r'=', header):
            missing += ' ' + attribute + '=' + quoteattr(uri)
    position = output.index('>')
    if output[position - 1] == '/':
        position -= 1
    return ('<?xml version="1.0" encoding="UTF-8"?>\n' + output[:position] + missing + output[position:]).encode('utf-8')


def patch_repeated(parent, tag, repetition, index):
    logical = 0
    for position, node in enumerate(list(parent)):
        if node.tag not in tag:
            continue
        count = int(node.get(repetition, 1))
        if logical <= index < logical + count:
            if count == 1:
                return node
            offset = index - logical
            parent.remove(node)
            inserts = []
            if offset:
                prefix = copy.deepcopy(node)
                prefix.set(repetition, str(offset))
                inserts.append(prefix)
            selected = copy.deepcopy(node)
            selected.attrib.pop(repetition, None)
            inserts.append(selected)
            if count - offset - 1:
                suffix = copy.deepcopy(node)
                suffix.set(repetition, str(count - offset - 1))
                inserts.append(suffix)
            for i, item in enumerate(inserts):
                parent.insert(position + i, item)
            return selected
        logical += count
    raise ValueError('The selected cell is outside this document’s stored sheet.')


def save_sheet(path, destination, edits):
    extension = path.suffix.lower()
    if extension in {'.csv', '.tsv'}:
        import csv
        raw = path.read_bytes()
        encoding = 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        text = raw.decode(encoding)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel_tab if extension == '.tsv' else csv.excel
        import io
        rows = list(csv.reader(io.StringIO(text), dialect))
        for edit in edits:
            row, column = int(edit['row']), int(edit['column'])
            while len(rows) <= row:
                rows.append([])
            while len(rows[row]) <= column:
                rows[row].append('')
            rows[row][column] = str(edit['value'])
        with destination.open('w', encoding=encoding, newline='') as stream:
            csv.writer(stream, dialect).writerows(rows)
        return
    if extension == '.ods':
        with zipfile.ZipFile(path) as archive:
            document = ET.fromstring(archive.read('content.xml'))
        tables = {t.get('{' + T + '}name'): t for t in document.iter('{' + T + '}table')}
        for edit in edits:
            table = tables.get(edit['sheet'])
            if table is None:
                raise ValueError('This worksheet no longer exists.')
            row = patch_repeated(table, {'{' + T + '}table-row'}, '{' + T + '}number-rows-repeated', int(edit['row']))
            cell = patch_repeated(row, {'{' + T + '}table-cell', '{' + T + '}covered-table-cell'}, '{' + T + '}number-columns-repeated', int(edit['column']))
            if cell.tag.endswith('covered-table-cell'):
                raise ValueError('Edit the first cell of a merged range instead.')
            for key in list(cell.attrib):
                if key.startswith('{' + O + '}') or key == '{' + T + '}formula':
                    cell.attrib.pop(key)
            for child in list(cell):
                cell.remove(child)
            cell.set('{' + O + '}value-type', 'string')
            ET.SubElement(cell, '{' + ODTEXT + '}p').text = str(edit['value'])
        with zipfile.ZipFile(path) as archive:
            original = archive.read('content.xml')
        rewrite_zip(path, destination, {'content.xml': serialize_xml(document, original)})
        return
    from openpyxl import Workbook, load_workbook
    if extension in {'.xls', '.xlsb'}:
        book = Workbook()
        book.remove(book.active)
        if extension == '.xls':
            import xlrd
            source = xlrd.open_workbook(str(path), on_demand=True)
            try:
                for original in source.sheets():
                    sheet = book.create_sheet(original.name)
                    for row in range(original.nrows):
                        sheet.append(original.row_values(row))
            finally:
                source.release_resources()
        else:
            from pyxlsb import open_workbook
            with open_workbook(str(path)) as source:
                for name in source.sheets:
                    sheet = book.create_sheet(name)
                    with source.get_sheet(name) as original:
                        for row in original.rows():
                            for cell in row:
                                if cell.v is not None:
                                    sheet.cell(cell.r + 1, cell.c + 1, cell.v)
    else:
        book = load_workbook(path, keep_vba=extension in {'.xlsm', '.xltm'}, keep_links=True)
    try:
        for edit in edits:
            value = edit['value']
            # Spreadsheet users can enter explicit formulas or numbers.
            if isinstance(value, str) and value and not value.startswith('='):
                try:
                    number = float(value)
                    if str(value).strip() == value and not (value.startswith('0') and len(value) > 1 and value[1].isdigit()):
                        value = int(number) if number.is_integer() else number
                except ValueError:
                    pass
            book[edit['sheet']].cell(int(edit['row']) + 1, int(edit['column']) + 1, value)
        book.save(destination)
    finally:
        book.close()


def install_document_editors(app, previews, root, context):
    def entry(identity):
        record = previews.files.get(identity)
        if not record or not Path(record['path']).is_file():
            raise HTTPException(404, 'This document is no longer available.')
        path = Path(record['path'])
        allowed_file(path)
        if path.stat().st_size > 100 * 1024 * 1024:
            raise HTTPException(413, 'The document editor supports files up to 100 MB. The original remains available to Vision.')
        return path

    @app.get('/api/document-edit/{identity}')
    async def read_document(identity: str, sheet: str = '', offset: int = 0, slide: int = 0):
        path = entry(identity)
        def read():
            extension = path.suffix.lower()
            if extension in SHEETS:
                result = spreadsheet(path, sheet, max(0, offset))
                result['converted'] = extension in {'.xls', '.xlsb'}
                if extension in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
                    from openpyxl import load_workbook
                    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                    try:
                        current = book[result['sheet']]
                        result['rows'] = [[str(v) if v is not None else '' for v in row] for row in current.iter_rows(min_row=result['offset'] + 1, max_row=min(current.max_row or 1, result['offset'] + 150), max_col=min(current.max_column or 1, 100), values_only=True)]
                    finally:
                        book.close()
            elif extension in IMAGE_EDIT:
                from PIL import Image, ImageOps
                image_plugins(path)
                with Image.open(path) as original:
                    image = ImageOps.exif_transpose(original)
                    result = {'kind': 'image', 'width': image.width, 'height': image.height, 'frames': getattr(original, 'n_frames', 1), 'format': original.format}
            elif extension in DOCUMENT_EDIT:
                zip_bound(path)
                part, document, paragraphs, namespace = xml_part(path, max(0, slide))
                result = {'kind': 'presentation' if extension == '.pptx' else 'document', 'part': part, 'paragraphs': [{'id': index, 'text': paragraph_text(paragraph, namespace)} for index, paragraph in enumerate(paragraphs)]}
                if extension == '.pptx':
                    result['preview'] = presentation(path, max(0, slide))
            else:
                raise ValueError('This format has no visual editing tool yet. Its preview and original file remain available to Vision.')
            result['revision'] = fingerprint(path)
            return result
        try:
            return await asyncio.to_thread(read)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, str(exc)) from exc

    @app.post('/api/document-edit/{identity}')
    async def save_document(identity: str, request: Request):
        path = entry(identity)
        data = await request.json()
        if len(json.dumps(data)) > 4 * 1024 * 1024:
            raise HTTPException(413, 'Save a smaller set of edits at once.')
        def write():
            if fingerprint(path) != data.get('revision'):
                raise HTTPException(409, 'This document changed. Reopen its editor before saving.')
            extension = path.suffix.lower()
            suffix = '.png' if extension in IMAGE_EDIT else '.xlsx' if extension in {'.xls', '.xlsb'} else extension
            folder = root / 'runtime/workspace-files' / secrets.token_hex(12)
            folder.mkdir(parents=True)
            destination = folder / (path.stem + '-edited' + suffix)
            try:
                if extension in SHEETS:
                    edits = data.get('cells', [])
                    if len(edits) > 10000:
                        raise ValueError('Save up to 10,000 cell changes at once.')
                    for edit in edits:
                        if not 0 <= int(edit['row']) < 1048576 or not 0 <= int(edit['column']) < 16384 or len(str(edit['value'])) > 32767:
                            raise ValueError('This cell is outside the supported spreadsheet limits.')
                    save_sheet(path, destination, edits)
                elif extension in IMAGE_EDIT:
                    from PIL import Image, ImageOps, ImageEnhance
                    image_plugins(path)
                    with Image.open(path) as original:
                        image = ImageOps.exif_transpose(original).convert('RGBA')
                        crop = data.get('crop', [0, 0, image.width, image.height])
                        x, y, width, height = map(int, crop)
                        if x < 0 or y < 0 or width < 1 or height < 1 or x + width > image.width or y + height > image.height:
                            raise ValueError('Keep the crop inside the original image.')
                        image = image.crop((x, y, x + width, y + height))
                        angle = int(data.get('rotation', 0))
                        if angle not in {0, 90, 180, 270}:
                            raise ValueError('Choose a quarter-turn rotation.')
                        if angle:
                            image = image.rotate(-angle, expand=True)
                        if data.get('flip'):
                            image = ImageOps.mirror(image)
                        dimensions = list(map(int, data.get('resize', image.size)))
                        if len(dimensions) != 2 or any(v < 1 or v > 16000 for v in dimensions) or dimensions[0] * dimensions[1] > 40000000:
                            raise ValueError('Choose dimensions up to 16,000 pixels per side and 40 megapixels total.')
                        image = image.resize(tuple(dimensions), Image.Resampling.LANCZOS)
                        for name, enhancer in [('brightness', ImageEnhance.Brightness), ('contrast', ImageEnhance.Contrast)]:
                            value = float(data.get(name, 1))
                            if not 0 <= value <= 3:
                                raise ValueError('Choose an adjustment between 0 and 3.')
                            image = enhancer(image).enhance(value)
                        image.save(destination, format='PNG')
                elif extension in DOCUMENT_EDIT:
                    changes = {}
                    for section in data.get('sections', []):
                        part, document, paragraphs, namespace = xml_part(path, max(0, int(section.get('slide', 0))))
                        for change in section.get('paragraphs', []):
                            index = int(change['id'])
                            text = str(change['text'])
                            if not 0 <= index < len(paragraphs) or len(text) > 100000:
                                raise ValueError('This paragraph cannot be edited.')
                            paragraph = paragraphs[index]
                            if namespace == ODTEXT:
                                for child in list(paragraph):
                                    paragraph.remove(child)
                                paragraph.text = text
                            else:
                                nodes = list(paragraph.iter('{' + namespace + '}t'))
                                if not nodes:
                                    run = ET.SubElement(paragraph, '{' + namespace + '}r')
                                    nodes = [ET.SubElement(run, '{' + namespace + '}t')]
                                nodes[0].text = text
                                nodes[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                                for node in nodes[1:]:
                                    node.text = ''
                        with zipfile.ZipFile(path) as archive:
                            original = archive.read(part)
                        changes[part] = serialize_xml(document, original)
                    rewrite_zip(path, destination, changes)
                else:
                    raise ValueError('This format is preview-only.')
                return destination
            except Exception:
                destination.unlink(missing_ok=True)
                raise
        try:
            destination = await asyncio.to_thread(write)
            result = await previews.register(destination, 'workspace', manual=True)
            if not result:
                raise ValueError('The edited copy could not be opened.')
            await context.add({**result, 'kind': 'file', 'path': str(destination)})
            await context.gateway.publish('settings.saved', {'text': 'Saved edited copy: ' + destination.name})
            return result
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, 'The edited copy could not be saved: ' + str(exc)) from exc
