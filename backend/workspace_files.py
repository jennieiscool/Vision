"""Read-only workspace previews. Opening a file never invokes its application."""
import asyncio
import base64
import csv
import hashlib
import io
import json
import math
import mimetypes
import posixpath
import re
import secrets
import shutil
import tarfile
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlsplit

from fastapi import HTTPException, Request, UploadFile
from fastapi.responses import FileResponse

APPLICATIONS = {'.exe', '.com', '.msi', '.msix', '.msixbundle', '.appx', '.appxbundle', '.dll', '.sys', '.scr', '.cpl', '.lnk', '.app', '.apk', '.dmg', '.pkg', '.deb', '.rpm', '.appimage', '.jar'}
SHEETS = {'.xlsx', '.xlsm', '.xltx', '.xltm', '.xls', '.xlsb', '.ods', '.csv', '.tsv'}
IMAGES = {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.avif', '.bmp', '.tif', '.tiff', '.ico', '.heic', '.heif'}
VIDEO = {'.mp4', '.webm', '.mkv', '.mov', '.avi', '.wmv', '.m4v', '.mpeg', '.mpg', '.3gp', '.ts', '.mts', '.m2ts'}
AUDIO = {'.mp3', '.wav', '.ogg', '.m4a', '.aac', '.flac', '.aiff', '.aif', '.wma', '.opus'}


def allowed_file(path):
    if path.suffix.lower() in APPLICATIONS:
        raise HTTPException(415, 'Applications and installers cannot be opened in the workspace.')
    with path.open('rb') as stream:
        signature = stream.read(4)
    if signature[:2] == b'MZ' or signature == b'\x7fELF' or signature in {b'\xcf\xfa\xed\xfe', b'\xfe\xed\xfa\xcf'}:
        raise HTTPException(415, 'Executable applications cannot be opened in the workspace.')


def cell(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return str(value)[:20000]


def zip_bound(path):
    with zipfile.ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > 512 * 1024 * 1024:
            raise ValueError('This document is too large for an inline preview.')


def spreadsheet(path, sheet, offset):
    extension = path.suffix.lower()
    limit, columns = 150, 100
    names, rows, total, selected = [], [], None, sheet
    if extension in {'.csv', '.tsv'}:
        with path.open('rb') as stream:
            sample = stream.read(8192)
        encoding = 'utf-16' if sample.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        try:
            dialect = csv.Sniffer().sniff(sample.decode(encoding, errors='replace'), delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel_tab if extension == '.tsv' else csv.excel
        with path.open(encoding=encoding, errors='replace', newline='') as stream:
            for index, row in enumerate(csv.reader(stream, dialect)):
                if index < offset:
                    continue
                rows.append([cell(v) for v in row[:columns]])
                if len(rows) > limit:
                    break
        names, selected = [path.stem], path.stem
    elif extension in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        from openpyxl import load_workbook
        zip_bound(path)
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            names = book.sheetnames
            selected = sheet if sheet in names else names[0]
            page = book[selected]
            total = page.max_row
            rows = [[cell(v) for v in row] for row in page.iter_rows(min_row=offset + 1, max_row=offset + limit + 1, max_col=min(page.max_column or 1, columns), values_only=True)]
        finally:
            book.close()
    elif extension == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            names = book.sheet_names()
            selected = sheet if sheet in names else names[0]
            page = book.sheet_by_name(selected)
            total = page.nrows
            rows = [[cell(v) for v in page.row_values(i)[:columns]] for i in range(offset, min(total, offset + limit + 1))]
        finally:
            book.release_resources()
    elif extension == '.xlsb':
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as book:
            names = book.sheets
            selected = sheet if sheet in names else names[0]
            with book.get_sheet(selected) as page:
                for index, row in enumerate(page.rows()):
                    if index >= offset:
                        rows.append([cell(v.v) for v in row[:columns]])
                    if len(rows) > limit:
                        break
    else:
        zip_bound(path)
        ns = {'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'}
        with zipfile.ZipFile(path) as archive:
            book = ET.fromstring(archive.read('content.xml'))
        sheets = book.findall('.//table:table', ns)
        table_ns = '{' + ns['table'] + '}'
        names = [table.get(table_ns + 'name', 'Sheet') for table in sheets]
        selected = sheet if sheet in names else names[0]
        index = 0
        for row in sheets[names.index(selected)].findall('table:table-row', ns):
            repeat = min(int(row.get(table_ns + 'number-rows-repeated', 1)), offset + limit + 1)
            values = []
            for item in row:
                if not item.tag.endswith('table-cell'):
                    continue
                values.extend([cell(''.join(item.itertext()))] * min(int(item.get(table_ns + 'number-columns-repeated', 1)), columns - len(values)))
                if len(values) >= columns:
                    break
            for _ in range(repeat):
                if index >= offset:
                    rows.append(values)
                index += 1
                if len(rows) > limit:
                    break
            if len(rows) > limit:
                break
    more = len(rows) > limit
    return {'kind': 'sheet', 'sheets': names, 'sheet': selected, 'rows': rows[:limit], 'offset': offset, 'more': more, 'total': total, 'column_limit': columns}


NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def relations(archive, part):
    filename = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
    if filename not in archive.namelist():
        return {}
    result = {}
    for item in ET.fromstring(archive.read(filename)):
        if item.get('TargetMode') == 'External':
            continue
        target = item.get('Target', '')
        result[item.get('Id')] = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part), target))
    return result


def embedded_image(archive, name):
    if not name or name not in archive.namelist() or archive.getinfo(name).file_size > 15 * 1024 * 1024:
        return None
    mime = mimetypes.guess_type(name)[0] or ''
    if not mime.startswith('image/'):
        return None
    return 'data:' + mime + ';base64,' + base64.b64encode(archive.read(name)).decode()


def document(path):
    # OpenXML is read directly. No Word installation or native application runs.
    zip_bound(path)
    blocks = []
    word = '{' + NS['w'] + '}'
    def paragraph_text(element):
        return ''.join(item.text or '' if item.tag == word + 't' else '\t' if item.tag == word + 'tab' else '\n' if item.tag == word + 'br' else '' for item in element.iter())
    with zipfile.ZipFile(path) as archive:
        doc = ET.fromstring(archive.read('word/document.xml'))
        refs = relations(archive, 'word/document.xml')
        body = doc.find('w:body', NS)
        for item in list(body)[:5000]:
            if item.tag == word + 'p':
                style = item.find('w:pPr/w:pStyle', NS)
                name = style.get(word + 'val', '') if style is not None else ''
                match = re.fullmatch(r'heading\s*([1-6])', name, re.I)
                level = int(match[1]) if match else 1 if name.lower() == 'title' else 0
                blocks.append({'kind': 'heading' if level else 'paragraph', 'level': level, 'text': paragraph_text(item)})
                for blip in item.findall('.//a:blip', NS):
                    image = embedded_image(archive, refs.get(blip.get('{' + NS['r'] + '}embed')))
                    if image:
                        blocks.append({'kind': 'image', 'url': image})
            elif item.tag == word + 'tbl':
                rows = [[ '\n'.join(paragraph_text(p) for p in value.findall('.//w:p', NS)) for value in row.findall('w:tc', NS)] for row in item.findall('w:tr', NS)[:1000]]
                blocks.append({'kind': 'table', 'rows': rows})
    return {'kind': 'document', 'blocks': blocks}


def presentation(path, index):
    zip_bound(path)
    with zipfile.ZipFile(path) as archive:
        deck = ET.fromstring(archive.read('ppt/presentation.xml'))
        size = deck.find('p:sldSz', NS)
        width, height = (int(size.get('cx')), int(size.get('cy'))) if size is not None else (9144000, 5143500)
        refs = relations(archive, 'ppt/presentation.xml')
        ids = deck.findall('p:sldIdLst/p:sldId', NS)
        if not ids:
            return {'kind': 'slides', 'count': 0, 'index': 0, 'ratio': width / height, 'shapes': []}
        index = min(index, len(ids) - 1)
        part = refs[ids[index].get('{' + NS['r'] + '}id')]
        slide = ET.fromstring(archive.read(part))
        refs = relations(archive, part)
        shapes = []
        for shape in slide.findall('.//p:sp', NS) + slide.findall('.//p:pic', NS) + slide.findall('.//p:graphicFrame', NS):
            position = shape.find('p:spPr/a:xfrm', NS)
            if position is None:
                position = shape.find('p:xfrm', NS)
            point = position.find('a:off', NS) if position is not None else None
            extent = position.find('a:ext', NS) if position is not None else None
            item = {'x': int(point.get('x', 0)) / width * 100 if point is not None else 5,
                    'y': int(point.get('y', 0)) / height * 100 if point is not None else 5 + len(shapes) * 12,
                    'w': int(extent.get('cx', width)) / width * 100 if extent is not None else 90,
                    'h': int(extent.get('cy', height)) / height * 100 if extent is not None else 15}
            blip = shape.find('.//a:blip', NS)
            table = shape.find('.//a:tbl', NS)
            if blip is not None:
                image = embedded_image(archive, refs.get(blip.get('{' + NS['r'] + '}embed')))
                if not image:
                    continue
                item.update(kind='image', url=image)
            elif table is not None:
                item.update(kind='table', rows=[[''.join(c.itertext()) for c in row.findall('a:tc', NS)] for row in table.findall('a:tr', NS)])
            else:
                item.update(kind='text', text='\n'.join(''.join(t.text or '' for t in paragraph.findall('.//a:t', NS)) for paragraph in shape.findall('.//a:p', NS)))
            shapes.append(item)
    return {'kind': 'slides', 'count': len(ids), 'index': index, 'ratio': width / height, 'shapes': shapes}


def read_preview(path, sheet, offset, page):
    if path.suffix.lower() in SHEETS | {'.docx', '.pptx', '.ods', '.odt', '.odp'} and path.stat().st_size > 100 * 1024 * 1024:
        raise ValueError('This document is too large for an inline preview. The original remains available to save.')
    extension = path.suffix.lower()
    if extension in SHEETS:
        return spreadsheet(path, sheet, offset)
    if extension == '.docx':
        return document(path)
    if extension == '.pptx':
        return presentation(path, page)
    if extension in {'.odt', '.odp'}:
        zip_bound(path)
        with zipfile.ZipFile(path) as archive:
            doc = ET.fromstring(archive.read('content.xml'))
        text_ns = '{urn:oasis:names:tc:opendocument:xmlns:text:1.0}'
        paragraphs = [''.join(item.itertext()) for item in doc.iter() if item.tag in {text_ns + 'p', text_ns + 'h'}]
        text = '\n'.join(paragraphs)
        return {'kind': 'text', 'text': text[:1000000], 'truncated': len(text) > 1000000}
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            return {'kind': 'archive', 'entries': [{'name': i.filename, 'size': i.file_size, 'folder': i.is_dir()} for i in entries[:2000]], 'total': len(entries)}
    if extension in {'.tar', '.gz', '.tgz', '.bz2', '.xz'} and tarfile.is_tarfile(path):
        entries = []
        with tarfile.open(path) as archive:
            for item in archive:
                entries.append({'name': item.name, 'size': item.size, 'folder': item.isdir()})
                if len(entries) >= 2000:
                    break
        return {'kind': 'archive', 'entries': entries, 'total': len(entries)}
    with path.open('rb') as stream:
        data = stream.read(1024 * 1024)
    encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
    try:
        text = data.decode(encoding)
        if '\x00' in text:
            raise UnicodeError('Binary file')
    except UnicodeError:
        return {'kind': 'binary', 'bytes': data[:512].hex(' '), 'size': path.stat().st_size}
    if extension == '.rtf':
        from striprtf.striprtf import rtf_to_text
        text = rtf_to_text(text)
    return {'kind': 'markdown' if extension in {'.md', '.markdown'} else 'text', 'text': text, 'truncated': path.stat().st_size > len(data)}


def install_workspace_files(app, previews, root, context):
    directory = root / 'runtime/workspace-files'
    directory.mkdir(exist_ok=True)
    conversion_lock = asyncio.Lock()

    def dismissed_path():
        return root / ('runtime/closed-previews-' + context.config.get('profile', 'default') + '.json')

    def dismissed():
        try:
            return json.loads(dismissed_path().read_text(encoding='utf-8'))
        except (OSError, ValueError):
            return {}

    app.state.closed_previews = dismissed

    @app.post('/api/workspace/previews/close')
    async def close_previews(request: Request):
        data = await request.json()
        identities = data.get('ids', [])
        if not isinstance(identities, list) or len(identities) > 2000:
            raise HTTPException(400, 'Choose the previews to close.')
        closed = dismissed()
        state = context.load()
        for identity in identities:
            if identity in previews.files:
                closed[identity] = previews.files[identity].get('modified', 0)
        state['items'] = [] if data.get('all') else [item for item in state['items'] if item['id'] not in identities and 'project-' + item['id'] not in identities]
        context.save(state)
        dismissed_path().write_text(json.dumps(closed), encoding='utf-8')
        await context.gateway.publish('workspace.context', context.public())
        return {'closed': identities, 'dismissed': closed}

    @app.post('/api/workspace/previews/reopen')
    async def reopen_preview(request: Request):
        closed = dismissed()
        closed.pop(str((await request.json()).get('id', '')), None)
        dismissed_path().write_text(json.dumps(closed), encoding='utf-8')
        return {'dismissed': closed}

    def entry(identity):
        record = previews.files.get(identity)
        if not record or not Path(record['path']).is_file():
            raise HTTPException(404, 'This file is no longer available.')
        path = Path(record['path'])
        allowed_file(path)
        return path

    @app.post('/api/workspace/open')
    async def open_local(request: Request):
        path = Path(str((await request.json()).get('path', '')))
        if not path.is_absolute() or not path.is_file():
            raise HTTPException(400, 'Choose a file to open.')
        allowed_file(path)
        result = await previews.register(path, 'workspace', manual=True)
        if not result:
            raise HTTPException(400, 'This file could not be opened.')
        await context.add({**result, 'kind': 'file', 'path': str(path)})
        return result

    @app.post('/api/workspace/files')
    async def upload(file: UploadFile):
        name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', Path(file.filename or 'file').name).strip(' .')[:180] or 'file'
        if Path(name).suffix.lower() in APPLICATIONS:
            raise HTTPException(415, 'Applications and installers cannot be opened in the workspace.')
        folder = directory / secrets.token_hex(12)
        folder.mkdir()
        path = folder / name
        try:
            size = 0
            with path.open('wb') as output:
                while chunk := await file.read(1024 * 1024):
                    size += len(chunk)
                    if size > 2 * 1024 * 1024 * 1024:
                        raise HTTPException(413, 'Use the desktop app to open files larger than 2 GB.')
                    output.write(chunk)
            allowed_file(path)
            result = await previews.register(path, 'workspace', manual=True)
            if not result:
                raise HTTPException(400, 'This file could not be opened.')
            await context.add({**result, 'kind': 'file', 'path': str(path)})
            return result
        except Exception:
            path.unlink(missing_ok=True)
            folder.rmdir()
            raise

    @app.get('/api/workspace/context')
    async def shared_context():
        return context.public()

    @app.post('/api/workspace/context')
    async def remove_context(request: Request):
        await context.remove(str((await request.json()).get('id', '')))
        return context.public()

    @app.post('/api/workspace/links')
    async def share_link(request: Request):
        url = str((await request.json()).get('url', '')).strip()
        parsed = urlsplit(url)
        if len(url) > 8000 or any(ord(c) < 32 for c in url) or parsed.scheme not in {'http', 'https'} or not parsed.hostname or parsed.username or parsed.password:
            raise HTTPException(400, 'Enter an http or https link.')
        if Path(parsed.path).suffix.lower() in APPLICATIONS:
            raise HTTPException(415, 'Application download links cannot be opened here.')
        item = {'id': 'url-' + hashlib.sha256(url.encode()).hexdigest()[:24], 'kind': 'url', 'name': parsed.hostname, 'url': url}
        await context.add(item)
        return item

    @app.get('/api/workspace/{identity}/view')
    async def view(identity: str, sheet: str = '', offset: int = 0, page: int = 0):
        path = entry(identity)
        try:
            data = await asyncio.to_thread(read_preview, path, sheet, max(0, offset), max(0, page))
            # Give the native agent a readable excerpt without needing HTTP cookies.
            text = ''
            if data['kind'] == 'sheet':
                text = 'Preview excerpt of ' + path.name + '\nWorksheet: ' + data['sheet'] + '\nRows start at ' + str(data['offset'] + 1) + '. This is a paged excerpt; the original workbook is available separately.\n\n' + '\n'.join('\t'.join(row) for row in data['rows'])
            elif data['kind'] == 'document':
                text = '\n\n'.join('\n'.join('\t'.join(row) for row in block['rows']) if block['kind'] == 'table' else block.get('text', '[Embedded image]') for block in data['blocks'])
            elif data['kind'] == 'slides':
                text = 'Preview excerpt: slide ' + str(data['index'] + 1) + ' of ' + str(data['count']) + '\n\n' + '\n\n'.join(item.get('text', '[Slide image or table]') for item in data['shapes'])
            elif data['kind'] in {'text', 'markdown'}:
                text = data['text']
            elif data['kind'] == 'archive':
                text = 'Archive directory listing; no files executed.\n\n' + '\n'.join(item['name'] for item in data['entries'])
            if text:
                readable = root / 'runtime/workspace-readable'
                readable.mkdir(exist_ok=True)
                (readable / (identity + '.txt')).write_text(text, encoding='utf-8')
            return data
        except Exception as exc:
            raise HTTPException(422, 'This file cannot be previewed: ' + str(exc)) from exc

    @app.get('/api/workspace/{identity}/image')
    async def image(identity: str):
        path = entry(identity)
        if path.suffix.lower() not in IMAGES:
            raise HTTPException(415, 'Not an image file.')
        target = directory / (identity + '.png')
        if not target.exists() or target.stat().st_mtime < path.stat().st_mtime:
            def convert():
                from PIL import Image, ImageOps
                if path.suffix.lower() in {".heic", ".heif"}:
                    try:
                        import pillow_heif
                        pillow_heif.register_heif_opener()
                    except ImportError:
                        pass
                with Image.open(path) as source:
                    source = ImageOps.exif_transpose(source)
                    source.thumbnail((6000, 6000))
                    source.convert('RGBA').save(target)
            try:
                await asyncio.to_thread(convert)
            except Exception as exc:
                ffmpeg = shutil.which('ffmpeg')
                if not ffmpeg:
                    raise HTTPException(422, 'This image format is unavailable on this computer.') from exc
                process = await asyncio.create_subprocess_exec(ffmpeg, '-y', '-v', 'error', '-i', str(path), '-frames:v', '1', str(target), stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.PIPE, creationflags=0x08000000)
                try:
                    await asyncio.wait_for(process.communicate(), 90)
                except (asyncio.TimeoutError, asyncio.CancelledError):
                    process.kill()
                    await process.wait()
                    target.unlink(missing_ok=True)
                    raise
                if process.returncode:
                    target.unlink(missing_ok=True)
                    raise HTTPException(422, 'This image format is unavailable on this computer.') from exc
        return FileResponse(target, media_type='image/png')

    @app.post('/api/workspace/{identity}/media')
    async def prepare_media(identity: str):
        path = entry(identity)
        if path.suffix.lower() not in VIDEO | AUDIO:
            raise HTTPException(415, 'Not a supported media file.')
        video = path.suffix.lower() in VIDEO
        target = directory / (identity + ('.mp4' if video else '.m4a'))
        async with conversion_lock:
            if not target.exists() or target.stat().st_mtime < path.stat().st_mtime:
                ffmpeg = shutil.which('ffmpeg')
                if not ffmpeg:
                    raise HTTPException(503, 'The local media converter is unavailable.')
                args = ['-y', '-v', 'error', '-i', str(path)]
                args += ['-map', '0:v:0', '-map', '0:a:0?', '-c:v', 'libx264', '-preset', 'veryfast', '-crf', '23', '-pix_fmt', 'yuv420p', '-vf', 'scale=trunc(iw/2)*2:trunc(ih/2)*2'] if video else ['-vn']
                args += ['-c:a', 'aac', '-b:a', '192k', '-movflags', '+faststart', str(target)]
                process = await asyncio.create_subprocess_exec(ffmpeg, *args, stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.PIPE, creationflags=0x08000000)
                try:
                    _, error = await asyncio.wait_for(process.communicate(), 1800)
                except (asyncio.CancelledError, asyncio.TimeoutError):
                    process.kill()
                    await process.wait()
                    target.unlink(missing_ok=True)
                    raise
                if process.returncode:
                    target.unlink(missing_ok=True)
                    raise HTTPException(422, 'The media could not be converted for preview.')
        return {'url': '/api/workspace/' + identity + '/media'}

    @app.get('/api/workspace/{identity}/media')
    async def media(identity: str):
        path = entry(identity)
        video = path.suffix.lower() in VIDEO
        target = directory / (identity + ('.mp4' if video else '.m4a'))
        if not target.is_file():
            raise HTTPException(404, 'Prepare this media preview first.')
        return FileResponse(target, media_type='video/mp4' if video else 'audio/mp4')
